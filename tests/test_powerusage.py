"""Behavior checks for parsing, merging and producing a fresh batch report."""

import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from main import InputError, combine, fill_missing_names, main, read_csv, write_excel


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates/2026_PowerUsageTable_template.xlsx"


class PowerUsageTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def source(self, name="PDU_2A1.csv", body=None):
        if body is None:
            body = ("06.10.2026;\nRack_No;11A\nRack_Name;Test Alpha\n"
                    "11A_PDU2_max;11A_PDU2_average\n0.35;0.24\n")
        path = self.folder / name
        path.write_text(body, encoding="utf-8-sig")
        return path

    def demo_records(self):
        readings = [r for p in sorted((ROOT / "demo/input").glob("*.csv"))
                    for r in read_csv(p)]
        return combine(readings)

    def test_demo_matches_by_rack_not_input_order(self):
        records = {r[0]: r for r in self.demo_records()}
        self.assertEqual(len(records), 3)
        self.assertEqual(records["2F_A11"][2:6], (0.32, 0.43, 0.24, 0.35))
        self.assertEqual(records["2F_B1"][2:6], (1.20, 1.55, 1.10, 1.40))
        self.assertEqual(records["2F_B2"][2:6], (0.95, 1.20, 0.88, 1.10))
        self.assertEqual(records["2F_B2"][1], "Тест B2")
        self.assertEqual(records["2F_A11"][-1], date(2026, 10, 6))

    def test_names_do_not_leak_to_the_next_rack(self):
        path = self.source(body=("06.10.2026;\nRack_Name;Alpha\n"
                                "11A_PDU2_max;11A_PDU2_average\n0.35;0.24\n"
                                "1B_PDU2_max;1B_PDU2_average\n1.4;1.1\n"))
        readings = read_csv(path)
        self.assertEqual(readings[0].name, "Alpha")
        self.assertIsNone(readings[1].name)
        with self.assertRaisesRegex(InputError, "Rack_Name is missing"):
            combine(readings)

    def test_date_and_name_follow_the_same_rack_across_blocks(self):
        path = self.source(body=("06.10.2026;\nRack_Name;Alpha\n"
                                "11A_PDU2_max;11A_PDU2_average\n0.35;0.24\n"
                                "2026-10-07;\n"
                                "A11_PDU2_average;A11_PDU2_max\n0;0\n"))
        readings = read_csv(path)
        self.assertEqual(readings[1].day, date(2026, 10, 7))
        self.assertEqual(readings[1].name, "Alpha")
        self.assertEqual(readings[1].average, 0)

    def test_decimal_comma_and_windows_encoding(self):
        path = self.source(body="06.10.2026;\nRack_Name;Тест\n11A_PDU2_average;11A_PDU2_max\n0,24;0,35\n")
        path.write_bytes(path.read_text(encoding="utf-8-sig").encode("cp1251"))
        reading = read_csv(path)[0]
        self.assertEqual((reading.name, reading.average, reading.maximum), ("Тест", 0.24, 0.35))

    def test_wrong_pdu_missing_date_invalid_number_and_incomplete_block(self):
        bad_cases = (
            ("06.10.2026;\n11A_PDU1_max;11A_PDU1_average\n0.35;0.24\n", "must contain PDU2"),
            ("11A_PDU2_max;11A_PDU2_average\n0.35;0.24\n", "date is missing"),
            ("06.10.2026;\n11A_PDU2_max;11A_PDU2_average\nNaN;0.24\n", "finite"),
            ("06.10.2026;\n11A_PDU2_max;11A_PDU2_average\n", "incomplete block"),
            ("31.02.2026;\n", "Invalid report date"),
            ("06.10.2026;\n11A_PDU2_max;11A_PDU2_average\n0.1;0.24\n", "Average exceeds Max"),
            ("06.10.2026;\nRack_No;1B\n11A_PDU2_max;11A_PDU2_average\n0.35;0.24\n", "Rack_No disagrees"),
        )
        for body, error in bad_cases:
            with self.subTest(error=error), self.assertRaisesRegex(InputError, error):
                read_csv(self.source(body=body))

    def test_identical_duplicate_is_ignored_and_missing_partner_is_optional(self):
        one = read_csv(self.source())[0]
        self.assertEqual(combine([one, one]), combine([one]))
        self.assertEqual(combine([one])[0][2:6], ("Неизвестно", "Неизвестно", 0.24, 0.35))
        with self.assertRaisesRegex(InputError, "both PDU1 and PDU2"):
            combine([one], require_both=True)

    def test_names_must_agree_between_sources(self):
        a = read_csv(self.source())[0]
        b = read_csv(self.source("PDU_2B1.csv", "06.10.2026;\nRack_Name;Other\n11A_PDU1_max;11A_PDU1_average\n0.43;0.32\n"))[0]
        with self.assertRaisesRegex(InputError, "Rack_Name differs"):
            combine([a, b])

    def test_excel_replaces_sample_values_and_preserves_styles_and_tariff(self):
        original_bytes = TEMPLATE.read_bytes()
        output = self.folder / "result.xlsx"
        records = self.demo_records()
        write_excel(TEMPLATE, output, records)
        before = load_workbook(TEMPLATE)
        after = load_workbook(output)
        try:
            self.assertEqual(before.sheetnames, after.sheetnames)
            for name in before.sheetnames:
                for row in before[name]:
                    for cell in row:
                        actual = after[name][cell.coordinate]
                        if name != "YaIDC" or cell.row == 1 or cell.column > 7:
                            self.assertEqual(actual.value, cell.value)
                        self.assertEqual(actual._style, cell._style)
            sheet = after["YaIDC"]
            populated = [row for row in sheet.iter_rows(min_row=2, max_col=7, values_only=True)
                         if any(value is not None for value in row)]
            self.assertEqual(len(populated), 3)
            self.assertEqual(sheet["A2"].value, "2F_A11")
            self.assertEqual(sheet["G2"].value, datetime(2026, 10, 6))
            self.assertEqual([sheet.cell(2, c).value for c in range(3, 7)], [0.32, 0.43, 0.24, 0.35])
            self.assertEqual(sheet["G2"].number_format, before["YaIDC"]["G2"].number_format)
            self.assertEqual(sheet.tables["FactPower"].ref, "A1:G4")
        finally:
            before.close()
            after.close()
        self.assertEqual(TEMPLATE.read_bytes(), original_bytes)

    def test_repeated_run_does_not_duplicate_a_rack_date(self):
        output = self.folder / "first.xlsx"
        repeated = self.folder / "second.xlsx"
        records = self.demo_records()
        write_excel(TEMPLATE, output, records)
        write_excel(output, repeated, records)
        workbook = load_workbook(repeated)
        try:
            sheet = workbook["YaIDC"]
            populated = [row for row in sheet.iter_rows(min_row=2, max_col=7, values_only=True)
                         if any(value is not None for value in row)]
            self.assertEqual(len(populated), 3)
            self.assertEqual(sheet.tables["FactPower"].ref, "A1:G4")
        finally:
            workbook.close()

    def test_reordered_columns_are_located_by_header(self):
        reordered = self.folder / "reordered.xlsx"
        workbook = load_workbook(TEMPLATE)
        sheet = workbook["YaIDC"]
        for row in sheet:
            row[2].value, row[4].value = row[4].value, row[2].value
        workbook.save(reordered)
        workbook.close()
        output = self.folder / "result.xlsx"
        write_excel(reordered, output, self.demo_records())
        workbook = load_workbook(output)
        try:
            sheet = workbook["YaIDC"]
            self.assertEqual(sheet["C2"].value, 0.24)
            self.assertEqual(sheet["E2"].value, 0.32)
        finally:
            workbook.close()

    def test_same_path_is_rejected_and_input_csv_without_names_cannot_pass(self):
        with self.assertRaisesRegex(InputError, "differ"):
            write_excel(TEMPLATE, TEMPLATE, self.demo_records())
        readings = [r for p in sorted((ROOT / "tests/fixtures/original").glob("*.csv")) for r in read_csv(p)]
        with self.assertRaisesRegex(InputError, "Rack_Name is missing"):
            combine(readings)

    def test_check_mode_does_not_create_an_output(self):
        output = self.folder / "absent.xlsx"
        self.assertEqual(main(["--input-dir", str(ROOT / "demo/input"), "--template", str(TEMPLATE),
                               "--output", str(output), "--check"]), 0)
        self.assertFalse(output.exists())

    def original_readings(self):
        return [r for p in sorted((ROOT / "tests/fixtures/original").glob("*.csv")) for r in read_csv(p)]

    def test_original_files_get_names_from_template_and_keep_their_values(self):
        workbook = load_workbook(TEMPLATE)
        try:
            records = combine(fill_missing_names(self.original_readings(), workbook))
        finally:
            workbook.close()
        self.assertEqual(records, [("2F_A11", "Selectel", 0.32, 0.43, 0.24, 0.35, date(2026, 10, 6))])
        output = self.folder / "original_result.xlsx"
        self.assertEqual(main(["--input-dir", str(ROOT / "tests/fixtures/original"), "--template", str(TEMPLATE),
                               "--output", str(output), "--rack-name-source", "template"]), 0)
        workbook = load_workbook(output)
        try:
            sheet = workbook["YaIDC"]
            populated = [row for row in sheet.iter_rows(min_row=2, max_col=7, values_only=True)
                         if any(value is not None for value in row)]
            self.assertEqual(len(populated), 1)
            self.assertEqual(sheet["B2"].value, "Selectel")
            self.assertEqual(sheet.tables["FactPower"].ref, "A1:G2")
        finally:
            workbook.close()

    def test_template_fallback_does_not_replace_existing_csv_names(self):
        reading = read_csv(self.source())[0]
        workbook = load_workbook(TEMPLATE)
        try:
            result = fill_missing_names([reading], workbook)
        finally:
            workbook.close()
        self.assertEqual(result[0].name, "Test Alpha")

    def test_template_name_lookup_marks_an_unmatched_rack_unknown(self):
        workbook = load_workbook(TEMPLATE)
        workbook["YaIDC"]["A2"] = "2F_A12"
        try:
            records = combine(fill_missing_names(self.original_readings(), workbook))
        finally:
            workbook.close()
        self.assertEqual(records[0][0], "2F_A11")
        self.assertEqual(records[0][1], "Неизвестно")
        self.assertEqual(records[0][2:6], (0.32, 0.43, 0.24, 0.35))

    def test_missing_name_is_written_to_excel_as_unknown(self):
        template = self.folder / "no_name.xlsx"
        workbook = load_workbook(TEMPLATE)
        workbook["YaIDC"]["B2"] = None
        workbook.save(template)
        workbook.close()
        output = self.folder / "unknown.xlsx"
        self.assertEqual(main(["--input-dir", str(ROOT / "tests/fixtures/original"), "--template", str(template),
                               "--output", str(output)]), 0)
        workbook = load_workbook(output)
        try:
            self.assertEqual(workbook["YaIDC"]["B2"].value, "Неизвестно")
            self.assertEqual(workbook["YaIDC"]["A2"].value, "2F_A11")
            self.assertEqual(workbook["YaIDC"]["C2"].value, 0.32)
        finally:
            workbook.close()

    def test_template_name_lookup_rejects_conflicting_names(self):
        workbook = load_workbook(TEMPLATE)
        workbook["YaIDC"]["A3"] = "2F_A11"
        workbook["YaIDC"]["B3"] = "Different Name"
        try:
            with self.assertRaisesRegex(InputError, "conflicting Rack_Name"):
                fill_missing_names(self.original_readings(), workbook)
        finally:
            workbook.close()

    def test_strict_csv_mode_still_requires_names_and_creates_no_file(self):
        output = self.folder / "not_created.xlsx"
        self.assertEqual(main(["--input-dir", str(ROOT / "tests/fixtures/original"), "--template", str(TEMPLATE),
                               "--output", str(output), "--rack-name-source", "strict-csv"]), 1)
        self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
