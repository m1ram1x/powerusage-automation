"""Batch reports across files, dates and racks, without inherited sample rows."""

import tempfile
import unittest
from dataclasses import replace
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from main import InputError, combine, discover_csv, main, read_csv, write_excel


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates/2026_PowerUsageTable_template.xlsx"


class BatchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.inputs = self.folder / "reports"
        self.inputs.mkdir()

    def tearDown(self):
        self.temp.cleanup()

    def report(self, filename, *, floor=2, line="B", day="2026-10-06",
               racks=("A11",), values=(0.32, 0.43), slot_only=False):
        pdu = 2 if line == "A" else 1
        source = f"PDU_{floor}{line}1"
        body = [] if slot_only else [source]
        body.extend([f"slot:/Services/ReportService/{source}", day])
        for rack in racks:
            body.extend([f"Rack_Name;Rack {rack}",
                         f"{rack}_PDU{pdu}_average;{rack}_PDU{pdu}_max",
                         f"{values[0]};{values[1]}"])
        path = self.inputs / filename
        path.write_text("\n".join(body) + "\n", encoding="utf-8-sig")
        return path

    def rows(self, workbook):
        return [row for row in workbook["YaIDC"].iter_rows(min_row=2, max_col=7,
                                                           values_only=True)
                if any(value is not None for value in row)]

    def test_source_is_read_from_content_of_renamed_reports(self):
        for filename in ("PDU_3A1_2026-10-06.csv", "any_report_name.csv"):
            with self.subTest(filename=filename):
                path = self.report(filename, floor=3, line="A", racks=("C7",))
                reading = read_csv(path)[0]
                self.assertEqual(reading.rack, "3F_C7")
                self.assertEqual(reading.pdu, 2)
                self.assertEqual(reading.source, filename)
                self.assertEqual(reading.day, date(2026, 10, 6))

    def test_source_can_be_read_from_slot_metadata(self):
        reading = read_csv(self.report("export.CSV", floor=4, line="B",
                                       racks=("B8",), slot_only=True))[0]
        self.assertEqual((reading.rack, reading.pdu), ("4F_B8", 1))

    def test_explicit_source_and_slot_must_agree(self):
        path = self.report("bad.csv")
        path.write_text(path.read_text(encoding="utf-8-sig").replace(
            "slot:/Services/ReportService/PDU_2B1", "slot:/Services/ReportService/PDU_3A1"),
            encoding="utf-8-sig")
        with self.assertRaises(InputError):
            read_csv(path)

    def test_multiple_explicit_source_sections_reset_date_and_names(self):
        path = self.report("combined.csv")
        with path.open("a", encoding="utf-8") as stream:
            stream.write("PDU_3A2\nslot:/Services/ReportService/PDU_3A2\n2026-10-07\n"
                         "A11_PDU2_average;A11_PDU2_max\n0.24;0.35\n")
        readings = read_csv(path)
        self.assertEqual(len(readings), 2)
        self.assertEqual((readings[0].rack, readings[0].pdu, readings[0].name),
                         ("2F_A11", 1, "Rack A11"))
        self.assertEqual((readings[1].rack, readings[1].pdu, readings[1].day, readings[1].name),
                         ("3F_A11", 2, date(2026, 10, 7), None))
        path.write_text(path.read_text(encoding="utf-8-sig").replace("2026-10-07\n", ""),
                        encoding="utf-8-sig")
        with self.assertRaisesRegex(InputError, "date is missing"):
            read_csv(path)

    def test_discovery_accepts_case_variants_and_ignores_non_csv_and_subfolders(self):
        files = [self.report("z_export.csv"), self.report("A_export.CSV"),
                 self.report("middle.CsV")]
        (self.inputs / "notes.txt").write_text("not a report", encoding="utf-8")
        nested = self.inputs / "nested.csv"
        nested.mkdir()
        (nested / "PDU_2B1.csv").write_text("not directly in input", encoding="utf-8")
        self.assertEqual(discover_csv(self.inputs), sorted(files))

    def test_many_files_merge_many_racks_and_dates_into_one_fresh_workbook(self):
        expected = {}
        for floor in (2, 3):
            for day_number in (6, 7):
                report_day = f"2026-10-{day_number:02}"
                for line in ("A", "B"):
                    racks = ("C3", "A1", "B2") if line == "A" else ("B2", "C3", "A1")
                    values = (0.24, 0.35) if line == "A" else (0.32, 0.43)
                    self.report(f"export_{floor}_{day_number}_{line}.CSV", floor=floor,
                                line=line, day=report_day, racks=racks, values=values)
                for rack in ("A1", "B2", "C3"):
                    expected[(f"{floor}F_{rack}", datetime(2026, 10, day_number))] = (
                        f"Rack {rack}", 0.32, 0.43, 0.24, 0.35)
        output = self.folder / "batch.xlsx"
        self.assertEqual(main(["--input-dir", str(self.inputs), "--template", str(TEMPLATE),
                               "--output", str(output)]), 0)
        workbook = load_workbook(output)
        try:
            rows = self.rows(workbook)
            self.assertEqual(len(rows), 12)
            actual = {(row[0], row[-1]): row[1:6] for row in rows}
            self.assertEqual(actual, expected)
            self.assertEqual(workbook["YaIDC"].tables["FactPower"].ref, "A1:G13")
        finally:
            workbook.close()

    def test_identical_readings_deduplicate_across_files_and_rows(self):
        reading = read_csv(self.report("first.csv"))[0]
        copied = replace(reading, source="copied.csv", row=125)
        self.assertEqual(combine([reading, copied]), combine([reading]))

    def test_conflicting_duplicate_reports_include_both_locations(self):
        reading = read_csv(self.report("first.csv"))[0]
        different = replace(reading, source="changed.csv", row=125, maximum=0.5)
        with self.assertRaises(InputError) as raised:
            combine([reading, different])
        message = str(raised.exception)
        self.assertIn("first.csv", message)
        self.assertIn("changed.csv", message)
        self.assertIn("125", message)

    def test_duplicate_names_cannot_silently_conflict(self):
        reading = read_csv(self.report("first.csv"))[0]
        different = replace(reading, name="Different Rack", source="changed.csv", row=125)
        with self.assertRaises(InputError) as raised:
            combine([reading, different])
        self.assertIn("first.csv", str(raised.exception))
        self.assertIn("changed.csv", str(raised.exception))

    def test_lone_b_source_has_pdu1_and_marks_only_pdu2_unknown(self):
        reading = read_csv(self.report("single.csv", line="B"))[0]
        self.assertEqual(combine([reading]), [
            ("2F_A11", "Rack A11", 0.32, 0.43, "Неизвестно", "Неизвестно",
             date(2026, 10, 6))])

    def test_zero_reading_is_retained_in_incomplete_pair(self):
        reading = read_csv(self.report("zero.csv", line="A", values=(0, 0)))[0]
        self.assertEqual(combine([reading])[0][2:6], ("Неизвестно", "Неизвестно", 0, 0))

    def test_cli_writes_single_source_and_optional_strict_mode_rejects_it(self):
        self.report("single.CSV", line="B")
        output = self.folder / "single.xlsx"
        argv = ["--input-dir", str(self.inputs), "--template", str(TEMPLATE),
                "--output", str(output)]
        self.assertEqual(main(argv + ["--require-both-pdus"]), 1)
        self.assertFalse(output.exists())
        self.assertEqual(main(argv), 0)
        workbook = load_workbook(output)
        try:
            self.assertEqual(self.rows(workbook), [
                ("2F_A11", "Rack A11", 0.32, 0.43, "Неизвестно", "Неизвестно",
                 datetime(2026, 10, 6))])
        finally:
            workbook.close()

    def test_default_uses_sample_template_only_for_structure_and_missing_name_is_unknown(self):
        output = self.folder / "real.xlsx"
        self.assertEqual(main(["--input-dir", str(ROOT / "tests/fixtures/original"), "--template", str(TEMPLATE),
                               "--output", str(output)]), 0)
        workbook = load_workbook(output)
        try:
            self.assertEqual(self.rows(workbook), [
                ("2F_A11", "Неизвестно", 0.32, 0.43, 0.24, 0.35,
                 datetime(2026, 10, 6))])
        finally:
            workbook.close()

    def test_known_csv_name_is_shared_when_partner_has_no_name(self):
        self.report("named.csv", line="B")
        nameless = self.report("nameless.csv", line="A", values=(0.24, 0.35))
        nameless.write_text(nameless.read_text(encoding="utf-8-sig").replace(
            "Rack_Name;Rack A11\n", ""), encoding="utf-8-sig")
        output = self.folder / "shared_name.xlsx"
        self.assertEqual(main(["--input-dir", str(self.inputs), "--template", str(TEMPLATE),
                               "--output", str(output)]), 0)
        workbook = load_workbook(output)
        try:
            self.assertEqual(self.rows(workbook), [
                ("2F_A11", "Rack A11", 0.32, 0.43, 0.24, 0.35,
                 datetime(2026, 10, 6))])
        finally:
            workbook.close()

    def test_output_replaces_sample_rows_but_preserves_other_cells_and_tariff(self):
        template = self.folder / "with_notes.xlsx"
        before = load_workbook(TEMPLATE)
        sheet = before["YaIDC"]
        sheet["H1"] = "Notes"
        sheet["H3"] = "Keep this note"
        sheet["H4"] = "=1+2"
        before.save(template)
        reading = read_csv(self.report("single.csv"))[0]
        output = self.folder / "fresh.xlsx"
        write_excel(template, output, combine([reading]))
        after = load_workbook(output)
        try:
            target = after["YaIDC"]
            self.assertEqual(len(self.rows(after)), 1)
            self.assertEqual(target.tables["FactPower"].ref, "A1:G2")
            self.assertEqual(target.tables["FactPower"].autoFilter.ref, "A1:G2")
            for row in range(3, 5):
                self.assertTrue(all(target.cell(row, col).value is None for col in range(1, 8)))
            for row in sheet:
                for cell in row:
                    actual = target[cell.coordinate]
                    if cell.row == 1 or cell.column > 7:
                        self.assertEqual(actual.value, cell.value)
                    # An unsaved unstyled cell has _style=None; loading may
                    # normalize it to the equivalent style array with id 0.
                    self.assertEqual(actual.style_id, cell.style_id)
            self.assertEqual(after.sheetnames, before.sheetnames)
            self.assertEqual(target.freeze_panes, sheet.freeze_panes)
            for key in sheet.column_dimensions:
                self.assertEqual(target.column_dimensions[key].width,
                                 sheet.column_dimensions[key].width)
            for row in before["Tariff"]:
                for cell in row:
                    self.assertEqual(after["Tariff"][cell.coordinate].value, cell.value)
                    self.assertEqual(after["Tariff"][cell.coordinate]._style, cell._style)
        finally:
            before.close()
            after.close()

    def test_repeated_batch_can_shrink_and_never_accumulates_old_rows(self):
        readings = read_csv(self.report("many.csv", racks=("A1", "B2", "C3", "D4", "E5")))
        first = self.folder / "first.xlsx"
        second = self.folder / "second.xlsx"
        write_excel(TEMPLATE, first, combine(readings))
        write_excel(first, second, combine([readings[0]]))
        workbook = load_workbook(second)
        try:
            self.assertEqual(len(self.rows(workbook)), 1)
            self.assertEqual(workbook["YaIDC"].tables["FactPower"].ref, "A1:G2")
        finally:
            workbook.close()

    def test_conflicting_input_does_not_overwrite_existing_output(self):
        self.report("first.csv", values=(0.32, 0.43))
        self.report("changed.csv", values=(0.32, 0.50))
        output = self.folder / "existing.xlsx"
        output.write_bytes(TEMPLATE.read_bytes())
        saved = output.read_bytes()
        self.assertEqual(main(["--input-dir", str(self.inputs), "--template", str(TEMPLATE),
                               "--output", str(output)]), 1)
        self.assertEqual(output.read_bytes(), saved)

    def test_truncated_rack_metadata_cannot_disappear_at_new_source(self):
        path = self.inputs / "truncated.csv"
        path.write_text("PDU_2B1;\n2026-10-06;\nRack_Name;Lost rack\n"
                        "PDU_3B1;\n2026-10-06;\nA1_PDU1_average;A1_PDU1_max\n0.2;0.3\n",
                        encoding="utf-8")
        with self.assertRaisesRegex(InputError, "incomplete rack metadata"):
            read_csv(path)

    def test_comma_report_with_semicolon_in_quoted_name(self):
        path = self.inputs / "quoted.csv"
        path.write_text('PDU_2B1,\n2026-10-06,\nRack_Name,"Rack; Alpha"\n'
                        'A11_PDU1_average,A11_PDU1_max\n0.32,0.43\n', encoding="utf-8")
        reading = read_csv(path)[0]
        self.assertEqual(reading.name, "Rack; Alpha")
        self.assertEqual((reading.average, reading.maximum), (0.32, 0.43))

    def test_growth_into_merged_cells_is_rejected_before_saving(self):
        template = self.folder / "merged.xlsx"
        workbook = load_workbook(TEMPLATE)
        workbook["YaIDC"].merge_cells("A5:B5")
        workbook.save(template)
        workbook.close()
        readings = read_csv(self.report("many.csv", racks=("A1", "B2", "C3", "D4", "E5")))
        output = self.folder / "not_created.xlsx"
        with self.assertRaisesRegex(InputError, "merged cells"):
            write_excel(template, output, combine(readings))
        self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
