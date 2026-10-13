"""Combine all CSV reports into a fresh Excel report using a template's layout."""

from __future__ import annotations

import argparse
import csv
import io
import json
import math
import re
import sys
import tempfile
from zipfile import BadZipFile
from copy import copy
from dataclasses import dataclass, replace
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from telegram_alerts import Alert, TelegramError, credentials, messages, send_alerts, send_excel


class InputError(ValueError):
    """An input cannot be interpreted without guessing."""


SOURCE = re.compile(r"PDU_(\d+)([AB])(\d+)", re.IGNORECASE)
SOURCE_PREFIX = re.compile(r"PDU_(\d+)([AB])(\d+)(?=$|[_\s(.\-])", re.IGNORECASE)
METRIC = re.compile(r"(.+)_PDU([12])_(max|average|avg)", re.IGNORECASE)
RACK_TOKEN = re.compile(r"(?:\d+F_[A-Z]+\d+|[A-Z]+\d+|\d+[A-Z]+)", re.IGNORECASE)
HEADERS = ("Rack_No.", "Rack_Name", "PDU1_avg", "PDU1_max",
           "PDU2_avg", "PDU2_max", "Date")
UNKNOWN = "Неизвестно"


def label(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def rack_number(raw: str, floor: str) -> str:
    """11A, A11 and 2F_A11 all identify 2F_A11 on floor 2."""
    raw = raw.strip().upper()
    full = re.fullmatch(r"(\d+)F_([A-Z]+)(\d+)", raw)
    if full:
        if int(full[1]) != int(floor):
            raise InputError(f"Rack {raw}: floor disagrees with the PDU source")
        return f"{int(full[1])}F_{full[2]}{int(full[3])}"
    forward = re.fullmatch(r"([A-Z]+)(\d+)", raw)
    backward = re.fullmatch(r"(\d+)([A-Z]+)", raw)
    if forward:
        letter, number = forward[1], forward[2]
    elif backward:
        letter, number = backward[2], backward[1]
    else:
        raise InputError(f"Unsupported rack identifier: {raw!r}")
    return f"{int(floor)}F_{letter}{int(number)}"


def parse_date(raw: str) -> date | None:
    for pattern, fmt in ((r"\d{2}\.\d{2}\.\d{4}", "%d.%m.%Y"),
                         (r"\d{4}-\d{2}-\d{2}", "%Y-%m-%d")):
        if re.fullmatch(pattern, raw):
            try:
                return datetime.strptime(raw, fmt).date()
            except ValueError as exc:
                raise InputError(f"Invalid report date: {raw!r}") from exc
    return None


@dataclass(frozen=True)
class Reading:
    rack: str
    name: str | None
    day: date
    pdu: int
    average: float
    maximum: float
    source: str
    row: int


def discover_csv(input_dir: Path) -> list[Path]:
    if not input_dir.is_dir():
        raise InputError(f"Input directory does not exist: {input_dir}")
    files = sorted(path for path in input_dir.iterdir()
                   if path.is_file() and path.suffix.lower() == ".csv")
    if not files:
        raise InputError(f"No CSV files found in {input_dir}")
    return files


def report_source(fields: list[str]):
    nonempty = [field for field in fields if field]
    if len(nonempty) == 1:
        match = SOURCE.fullmatch(nonempty[0])
        if match:
            return match, "header"
    if fields and fields[0].lower().startswith("slot:"):
        match = SOURCE.fullmatch(fields[0].rstrip("/").rsplit("/", 1)[-1])
        if match:
            return match, "slot"
    return None


def csv_delimiter(text: str) -> str:
    # Recognize a metric-header row, so a ';' inside a quoted rack name does
    # not cause a comma-separated export to be read with the wrong separator.
    for line in text.splitlines():
        if "_pdu" not in line.lower():
            continue
        for delimiter in (";", "\t", ","):
            try:
                fields = next(csv.reader([line], delimiter=delimiter, strict=True))
            except csv.Error:
                continue
            nonempty = [field.strip() for field in fields if field.strip()]
            matches = [METRIC.fullmatch(field) for field in nonempty]
            if nonempty and all(match and RACK_TOKEN.fullmatch(match[1]) for match in matches):
                return delimiter
    first = next((line for line in text.splitlines() if line.strip()), "")
    return ";" if ";" in first else "\t" if "\t" in first else ","


def audit_csv(path: Path) -> list[Alert]:
    """Scan dates and raw Max independently, including files the parser rejects.

    Keep checking later blocks after a bad date: an invalid report must still
    produce useful alarms. The strict parser remains responsible for Excel.
    """
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1251")
    rows = csv.reader(io.StringIO(text), delimiter=csv_delimiter(text), strict=True)
    identity = SOURCE_PREFIX.match(path.stem)
    source = identity.group(0).upper() if identity else path.stem
    # Resolve arbitrary export filenames before reporting bad preceding metadata.
    for probe in csv.reader(io.StringIO(text), delimiter=csv_delimiter(text), strict=True):
        found = report_source([value.strip() for value in probe])
        if found:
            source = found[0].group(0).upper()
            break
        if any(METRIC.fullmatch(value.strip()) for value in probe):
            break
    alerts = []
    day = None
    date_problem = False
    pending = None

    def add(row_no, description):
        alerts.append(Alert(source, path.name, row_no, description))

    for row in rows:
        fields = [value.strip() for value in row]
        nonempty = [value for value in fields if value]
        if not nonempty:
            continue
        found = report_source(fields)
        if found:
            if pending is not None:
                add(pending[1], "Отсутствует строка значений Max после заголовка.")
                pending = None
            identifier, origin = found
            if origin == "header" or identifier.group(0).upper() != source:
                day = None
                date_problem = False
            source = identifier.group(0).upper()
            continue
        if pending is not None:
            columns, _ = pending
            for index, rack in columns:
                raw_max = fields[index] if index < len(fields) else ""
                try:
                    maximum = float(raw_max.replace(",", "."))
                    if not math.isfinite(maximum):
                        raise ValueError()
                except ValueError:
                    add(rows.line_num, f"{rack}: Max отсутствует или не является конечным числом.")
                    continue
                if maximum < 0.1 or maximum > 4:
                    bound = "ниже 0,1" if maximum < 0.1 else "выше 4"
                    when = day.isoformat() if day else "дата неизвестна/некорректна"
                    add(rows.line_num, f"Шкаф: {rack}\nДата: {when}\nMax: {maximum:g} кВт")
            pending = None
            continue
        first = fields[0]
        if first.lower().startswith("slot:"):
            continue
        if len(nonempty) == 1 and re.match(r"^\d{1,4}[./-]\d", first):
            try:
                parsed = parse_date(first)
            except InputError:
                parsed = None
            if parsed is None:
                add(rows.line_num, f"Некорректная дата «{first[:100]}»: нужна существующая дата ДД.ММ.ГГГГ или ГГГГ-ММ-ДД.")
                day = None
                date_problem = True
            else:
                day = parsed
                date_problem = False
            continue
        matches = [METRIC.fullmatch(value) if value else None for value in fields]
        if any(matches):
            if day is None and not date_problem:
                add(rows.line_num, "Перед показаниями отсутствует корректная дата (ДД.ММ.ГГГГ или ГГГГ-ММ-ДД).")
                date_problem = True
            columns = [(index, match[1]) for index, match in enumerate(matches)
                       if match and match[3].lower() == "max"]
            if not columns:
                add(rows.line_num, "В блоке показаний отсутствует столбец Max; проверить порог невозможно.")
            pending = columns, rows.line_num
    if pending is not None:
        add(pending[1], "Отсутствует строка значений Max после заголовка.")
    return alerts


def read_csv(path: Path) -> list[Reading]:
    """Names are bound to racks using only metadata preceding their readings."""
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1251")
    delimiter = csv_delimiter(text)
    identity = None
    for probe in csv.reader(io.StringIO(text), delimiter=delimiter, strict=True):
        found = report_source([field.strip() for field in probe])
        if found:
            identity = found[0]
            break
        if any(METRIC.fullmatch(field.strip()) for field in probe if field.strip()):
            break
    if identity is None:
        identity = SOURCE_PREFIX.match(path.stem)
    if identity is None:
        raise InputError(f"{path.name}: missing PDU_<floor><A|B><device> source in contents or filename")
    current_source = identity.group(0).upper()
    floor, line = identity[1], identity[2].upper()
    expected_pdu = 2 if line == "A" else 1
    source_locked = False
    rows = csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
    report_day = None
    declared_rack = None
    pending_name = None
    names: dict[str, str] = {}
    columns = None
    readings = []
    for row in rows:
        row_no = rows.line_num
        fields = [v.strip() for v in row]
        if not any(fields):
            continue
        nonempty = [v for v in fields if v]
        if columns is not None:
            if report_source(fields):
                raise InputError(f"{path.name}:{row_no}: incomplete metric block before a new source")
            if len(fields) < len(columns):
                raise InputError(f"{path.name}:{row_no}: too few metric values")
            pairs = {}
            for col, spec in enumerate(columns):
                if spec is None:
                    if fields[col]:
                        raise InputError(f"{path.name}:{row_no}: unexpected extra value")
                    continue
                rack, metric = spec
                try:
                    value = float(fields[col].replace(",", "."))
                except ValueError as exc:
                    raise InputError(f"{path.name}:{row_no}: invalid number {fields[col]!r}") from exc
                if not math.isfinite(value):
                    raise InputError(f"{path.name}:{row_no}: number must be finite")
                pairs.setdefault(rack, {})[metric] = value
            if any(fields[len(columns):]):
                raise InputError(f"{path.name}:{row_no}: unexpected extra columns")
            for rack, values in pairs.items():
                if set(values) != {"max", "avg"}:
                    raise InputError(f"{path.name}:{row_no}: Average/Max pair missing for {rack}")
                if values["avg"] > values["max"]:
                    raise InputError(f"{path.name}:{row_no}: Average exceeds Max for {rack}")
                readings.append(Reading(rack, names.get(rack), report_day,
                                        expected_pdu, values["avg"], values["max"],
                                        path.name, row_no))
            columns = None
            continue
        first = fields[0]
        found = report_source(fields)
        if found:
            identity, origin = found
            identifier = identity.group(0).upper()
            if origin == "slot" and source_locked:
                if identifier != current_source:
                    raise InputError(f"{path.name}:{row_no}: source header and slot disagree")
                continue
            if declared_rack is not None or pending_name is not None:
                raise InputError(f"{path.name}:{row_no}: incomplete rack metadata before a new source")
            current_source = identifier
            floor, line = identity[1], identity[2].upper()
            expected_pdu = 2 if line == "A" else 1
            source_locked = True
            report_day = declared_rack = pending_name = None
            names = {}
            continue
        if first.lower().startswith("slot:"):
            continue
        if len(nonempty) == 1:
            try:
                new_day = parse_date(first)
            except InputError as exc:
                raise InputError(f"{path.name}:{row_no}: {exc}") from exc
            if new_day is not None:
                report_day = new_day
                continue
        key = label(first)
        if key in {"rackno", "racknumber", "rackname"}:
            if len(nonempty) != 2 or len(fields) < 2 or not fields[1]:
                raise InputError(f"{path.name}:{row_no}: metadata needs a label and value")
            if key == "rackname":
                if pending_name is not None:
                    raise InputError(f"{path.name}:{row_no}: duplicate Rack_Name before metrics")
                pending_name = fields[1]
            else:
                if declared_rack is not None:
                    raise InputError(f"{path.name}:{row_no}: duplicate Rack_No before metrics")
                declared_rack = rack_number(fields[1], floor)
            continue
        matches = [METRIC.fullmatch(v) if v else None for v in fields]
        if not any(matches) or any(v and m is None for v, m in zip(fields, matches)):
            raise InputError(f"{path.name}:{row_no}: unsupported row {fields!r}")
        if report_day is None:
            raise InputError(f"{path.name}:{row_no}: report date is missing before metrics")
        columns = []
        seen = set()
        for match in matches:
            if match is None:
                columns.append(None)
                continue
            if int(match[2]) != expected_pdu:
                raise InputError(f"{path.name}:{row_no}: line {line} must contain PDU{expected_pdu}")
            spec = (rack_number(match[1], floor),
                    "avg" if match[3].lower() in {"avg", "average"} else "max")
            if spec in seen:
                raise InputError(f"{path.name}:{row_no}: duplicate metric {spec}")
            seen.add(spec)
            columns.append(spec)
        racks = {spec[0] for spec in columns if spec is not None}
        if declared_rack is not None and racks != {declared_rack}:
            raise InputError(f"{path.name}:{row_no}: Rack_No disagrees with metric headers")
        if pending_name is not None:
            if len(racks) != 1:
                raise InputError(f"{path.name}:{row_no}: Rack_Name is ambiguous across multiple racks")
            names[next(iter(racks))] = pending_name
        declared_rack = pending_name = None
    if columns is not None or pending_name is not None or declared_rack is not None:
        raise InputError(f"{path.name}: incomplete block at end of file")
    if not readings:
        raise InputError(f"{path.name}: no PDU readings found")
    return readings


def combine(readings: list[Reading], *, require_both: bool = False) -> list[tuple]:
    grouped = {}
    for item in readings:
        if item.name is None:
            raise InputError(f"{item.source}:{item.row}: Rack_Name is missing for {item.rack}; "
                             "add rack metadata in the source CSV before its metrics")
        key = item.rack, item.day
        pdus = grouped.setdefault(key, {})
        if item.pdu in pdus:
            previous = pdus[item.pdu]
            locations = f"{previous.source}:{previous.row} and {item.source}:{item.row}"
            if (previous.average, previous.maximum) != (item.average, item.maximum):
                raise InputError(f"Conflicting PDU{item.pdu} readings for {item.rack} on {item.day}: {locations}")
            names = {previous.name, item.name} - {UNKNOWN}
            if len(names) > 1:
                raise InputError(f"Rack_Name differs for {item.rack} on {item.day}: {locations}")
            pdus[item.pdu] = replace(previous, name=next(iter(names), UNKNOWN))
            continue
        pdus[item.pdu] = item
    result = []
    for (rack, day), pdus in sorted(grouped.items()):
        if require_both and set(pdus) != {1, 2}:
            raise InputError(f"{rack} on {day}: both PDU1 and PDU2 are required")
        names = {item.name for item in pdus.values()} - {UNKNOWN}
        if len(names) > 1:
            raise InputError(f"{rack} on {day}: Rack_Name differs between the two sources")
        one, two = pdus.get(1), pdus.get(2)
        result.append((rack, next(iter(names), UNKNOWN),
                       one.average if one else UNKNOWN, one.maximum if one else UNKNOWN,
                       two.average if two else UNKNOWN, two.maximum if two else UNKNOWN, day))
    return result


def table_layout(workbook):
    matches = []
    required = {label(h) for h in HEADERS}
    for sheet in workbook:
        for row_no, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20)), 1):
            mapping = {}
            duplicate_headers = set()
            for cell in row:
                key = label(cell.value)
                if key in required:
                    if key in mapping:
                        duplicate_headers.add(key)
                    mapping[key] = cell.column
            if required.issubset(mapping):
                if duplicate_headers:
                    raise InputError(f"{sheet.title}: duplicate columns {sorted(duplicate_headers)}")
                matches.append((sheet, row_no, [mapping[label(h)] for h in HEADERS]))
    if len(matches) != 1:
        raise InputError("Template must contain exactly one table with the seven required headers")
    return matches[0]


def fill_missing_names(readings: list[Reading], workbook) -> list[Reading]:
    """Look up missing names in the template; label unresolved names explicitly."""
    needed = {item.rack for item in readings if item.name is None}
    if not needed:
        return readings
    sheet, header_row, cols = table_layout(workbook)
    names = {}
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        raw_rack, raw_name = values[cols[0] - 1], values[cols[1] - 1]
        rack = str(raw_rack).strip().upper()
        if rack not in needed or raw_name is None:
            continue
        if not isinstance(raw_name, str) or not raw_name.strip():
            continue
        name = raw_name.strip()
        if rack in names and names[rack] != name:
            raise InputError(f"Template has conflicting Rack_Name values for {rack}")
        names[rack] = name
    return [replace(item, name=names.get(item.rack, UNKNOWN)) if item.name is None else item
            for item in readings]


def write_excel(template: Path, output: Path, records: list[tuple]) -> int:
    """Replace sample table values with this batch; preserve the workbook layout."""
    if template.resolve() == output.resolve():
        raise InputError("Output must differ from the original template")
    if not records:
        raise InputError("No records to write")
    workbook = load_workbook(template)
    try:
        sheet, header_row, cols = table_layout(workbook)
        last_row = header_row + len(records)
        if last_row > 1_048_576:
            raise InputError("The combined report exceeds Excel's row limit")
        target_tables = []
        for table in sheet.tables.values():
            left, top, right, bottom = range_boundaries(table.ref)
            if top == header_row and left <= min(cols) and right >= max(cols):
                if table.totalsRowShown or table.totalsRowCount:
                    raise InputError("Templates with a table totals row are not supported")
                target_tables.append((table, left, right, bottom))
        if len(target_tables) > 1:
            raise InputError("Template has overlapping report tables")
        body_bottom = target_tables[0][3] if target_tables else sheet.max_row
        clear_cols = range(target_tables[0][1], target_tables[0][2] + 1) if target_tables else cols
        first_data_row = header_row + 1
        for merged in sheet.merged_cells.ranges:
            if (merged.min_row <= max(body_bottom, last_row)
                    and merged.max_row >= first_data_row
                    and merged.min_col <= max(clear_cols) and merged.max_col >= min(clear_cols)):
                raise InputError(f"Report area overlaps merged cells: {merged}")
        styles = {col: copy(sheet.cell(first_data_row, col)._style) for col in cols}
        row_height = sheet.row_dimensions[first_data_row].height
        # Check before writing so growth cannot erase unrelated cells or tables.
        for row_no in range(body_bottom + 1, last_row + 1):
            if any(sheet.cell(row_no, col).value is not None for col in clear_cols):
                raise InputError(f"Report growth would overwrite unrelated data at {sheet.title}!row {row_no}")
        for table in sheet.tables.values():
            if target_tables and table is target_tables[0][0]:
                continue
            left, top, right, bottom = range_boundaries(table.ref)
            if top <= last_row and bottom >= first_data_row and left <= max(cols) and right >= min(cols):
                raise InputError(f"Report overlaps another table: {table.name}")
        for row_no in range(first_data_row, body_bottom + 1):
            for col in clear_cols:
                sheet.cell(row_no, col).value = None
        for row_no, record in enumerate(records, first_data_row):
            if row_no > body_bottom:
                for col in cols:
                    sheet.cell(row_no, col)._style = copy(styles[col])
                sheet.row_dimensions[row_no].height = row_height
            for col, value in zip(cols, record):
                cell = sheet.cell(row_no, col)
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = "s"
            date_cell = sheet.cell(row_no, cols[-1])
            if date_cell.number_format == "General":
                date_cell.number_format = "yyyy-mm-dd"
        # Fit the longer placeholder and CSV names without restyling other columns.
        name_letter = get_column_letter(cols[1])
        current_dimension = sheet.column_dimensions.get(name_letter)
        current_width = (current_dimension.width if current_dimension else
                         sheet.sheet_format.defaultColWidth or 8.43)
        desired_width = min(40.0, max(13.0, max(len(str(record[1])) for record in records) + 3.0))
        sheet.column_dimensions[name_letter].width = max(current_width, desired_width)
        for row_no, record in enumerate(records, first_data_row):
            if len(str(record[1])) > desired_width - 3:
                cell = sheet.cell(row_no, cols[1])
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                cell.alignment = alignment
                lines = math.ceil(len(str(record[1])) / (desired_width - 3))
                height = (cell.font.sz or 11) * 1.4 * lines
                sheet.row_dimensions[row_no].height = max(sheet.row_dimensions[row_no].height or 15, height)
        for table, left, right, _ in target_tables:
            old_ref = table.ref
            table.ref = f"{get_column_letter(left)}{header_row}:{get_column_letter(right)}{last_row}"
            if table.autoFilter is not None:
                table.autoFilter.ref = table.ref
            if sheet.auto_filter.ref == old_ref:
                sheet.auto_filter.ref = table.ref
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as stream:
            temp_path = Path(stream.name)
        try:
            workbook.save(temp_path)
            temp_path.replace(output)
        finally:
            temp_path.unlink(missing_ok=True)
        return len(records)
    finally:
        workbook.close()


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=Path("input"))
    parser.add_argument("--template", type=Path,
                        default=Path("templates/2026_PowerUsageTable_template.xlsx"))
    parser.add_argument("--output", type=Path,
                        default=Path("output/2026_PowerUsageTable.xlsx"))
    parser.add_argument("--check", action="store_true", help="Validate and print records without writing Excel")
    parser.add_argument("--rack-name-source", choices=("csv", "template", "strict-csv"), default="csv",
                        help="csv: CSV name or Неизвестно; template: allow template lookup; strict-csv: require CSV name")
    parser.add_argument("--require-both-pdus", action="store_true",
                        help="Reject incomplete pairs instead of marking missing PDU cells Неизвестно")
    parser.add_argument("--alerts-preview", action="store_true",
                        help="Print date/Max warnings without sending anything")
    parser.add_argument("--telegram", action="store_true",
                        help="Send date/Max warnings to the configured Telegram chat")
    parser.add_argument("--telegram-env", type=Path,
                        help="Local KEY=value credential file; otherwise use environment variables")
    arguments = sys.argv[1:] if argv is None else list(argv)
    args = parser.parse_args(arguments)
    auto_output_name = not any(arg == "--output" or arg.startswith("--output=")
                               for arg in arguments)
    if args.telegram and (args.check or args.alerts_preview):
        parser.error("--check/--alerts-preview cannot be combined with --telegram")
    alerts = []
    telegram_config = None
    status = 0
    try:
        if args.telegram:
            telegram_config = credentials(args.telegram_env)
        files = discover_csv(args.input_dir)
        if args.output.suffix.lower() != ".xlsx":
            raise InputError("Output filename must end with .xlsx")
        if args.output.resolve() in {p.resolve() for p in files}:
            raise InputError("Output must differ from all input files")
        if args.telegram_env and args.output.resolve() == args.telegram_env.resolve():
            raise InputError("Output must differ from the Telegram credential file")
        readings = []
        input_errors = []
        for path in files:
            try:
                alerts.extend(audit_csv(path))
            except (OSError, csv.Error, UnicodeError) as exc:
                input_errors.append(f"{path.name}: cannot audit CSV ({type(exc).__name__})")
                continue
            try:
                readings.extend(read_csv(path))
            except (InputError, OSError, csv.Error, UnicodeError) as exc:
                message = str(exc)
                input_errors.append(message if message.startswith(path.name + ":") else f"{path.name}: {message}")
        if input_errors:
            raise InputError("\n".join(input_errors))
        with_template = load_workbook(args.template, read_only=True)
        try:
            table_layout(with_template)
            if args.rack_name_source == "template":
                readings = fill_missing_names(readings, with_template)
            elif args.rack_name_source == "csv":
                readings = [replace(item, name=UNKNOWN) if item.name is None else item
                            for item in readings]
            records = combine(readings, require_both=args.require_both_pdus)
            if auto_output_name:
                dates = sorted({record[-1] for record in records})

                if len(dates) == 1:
                    date_label = dates[0].isoformat()
                else:
                    date_label = (
                        f"{dates[0].isoformat()}_to_{dates[-1].isoformat()}"
                    )

                args.output = (
                        args.output.parent
                        / f"PowerUsage_{date_label}.xlsx"
                )
        finally:
            with_template.close()
        if args.check:
            print(json.dumps([dict(zip(HEADERS, row)) for row in records],
                             ensure_ascii=False, default=str, indent=2))
        else:
            count = write_excel(args.template, args.output, records)
            print(f"Read {len(files)} CSV file(s), {len(readings)} PDU reading(s).")
            print(f"Saved {count} rack/date record(s): {args.output.resolve()}")
    except (InputError, OSError, csv.Error, UnicodeError, BadZipFile, InvalidFileException) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        status = 1
    except TelegramError as exc:
        print(f"TELEGRAM: {exc}", file=sys.stderr)
        return 2
    texts = messages(alerts)
    if args.alerts_preview:
        for text in texts:
            print(text)
        if not texts:
            print("No date/Max alerts.")
    elif not args.check:
        # Never overwrite inputs, templates or a chosen credential file with logs.
        alert_path = args.output.parent / "alerts.txt"
        protected = {args.template.resolve(), args.output.resolve()}
        if args.telegram_env:
            protected.add(args.telegram_env.resolve())
        try:
            if alert_path.resolve() in protected:
                raise OSError("Alert log path overlaps a protected file")
            alert_path.parent.mkdir(parents=True, exist_ok=True)
            alert_path.write_text("\n\n".join(texts) + ("\n" if texts else ""), encoding="utf-8")
            if texts:
                print(f"Date/Max alerts: {len(alerts)}; saved to {alert_path.resolve()}")
        except OSError:
            print("ERROR: cannot save the local alert report", file=sys.stderr)
            status = status or 1
    if telegram_config:
        try:
            sent = send_alerts(alerts, *telegram_config)
            print(f"Telegram messages sent: {sent}")
            if status == 0:
                send_excel(args.output, *telegram_config)
                print(f"Excel отправлен в Telegram")
        except TelegramError as exc:
            print(f"TELEGRAM: {exc}; see local alerts.txt", file=sys.stderr)
            return status or 2
    return status


if __name__ == "__main__":
    raise SystemExit(main())
